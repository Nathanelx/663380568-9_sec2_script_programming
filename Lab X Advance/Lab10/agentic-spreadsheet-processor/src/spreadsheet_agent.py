# agentic-spreadsheet-processor/src/spreadsheet_agent.py
import datetime
import json

import openpyxl
from openpyxl.styles import Font, PatternFill, Color
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, Series
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, IconSetRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter, range_boundaries

from .utils import setup_logging  # Use utils for consistent logging setup

logger = setup_logging(__name__)


class SpreadsheetAgent:
    """
    An agent that performs a series of operations on an Excel workbook
    based on a provided configuration.
    """
    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []
        self.start_time = datetime.datetime.now()
        logger.info("Spreadsheet Agent initialized.")

    def _log_audit(self, status, message, task_type="N/A", details=None):
        """Adds an entry to the audit log."""
        log_entry = {
            "timestamp": datetime.datetime.now().isoformat(),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details if details else {}
        }
        self.audit_log.append(log_entry)
        if status in ("ERROR", "CRITICAL"):
            logger.error(f"AUDIT LOG - {status}: {message} | Task: {task_type} | Details: {details}")
        else:
            logger.info(f"AUDIT LOG - {status}: {message} | Task: {task_type}")

    def _load_workbook(self):
        """Loads the input workbook."""
        input_file = self.config["input_file"]
        try:
            self.workbook = openpyxl.load_workbook(input_file)
            self._log_audit("SUCCESS", f"Workbook '{input_file}' loaded.", "load_workbook")
            return True
        except FileNotFoundError:
            self._log_audit("ERROR", f"Input file not found: {input_file}", "load_workbook")
            return False
        except Exception as e:
            self._log_audit("ERROR", f"Failed to load workbook '{input_file}': {e}", "load_workbook")
            return False

    def _get_sheet(self, sheet_name):
        """Safely gets a worksheet from the workbook (creates it if missing)."""
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        self._log_audit("WARNING", f"Sheet '{sheet_name}' not found. Creating new sheet.", "get_sheet")
        return self.workbook.create_sheet(sheet_name)

    def _process_copy_data_task(self, task):
        """Copies data from a source sheet to a destination sheet."""
        src_sheet_name = task.get("source_sheet")
        dest_sheet_name = task.get("dest_sheet")

        if not src_sheet_name or not dest_sheet_name:
            self._log_audit("ERROR", "Missing source_sheet or dest_sheet for copy_data task.", "copy_data", task)
            return

        src_sheet = self._get_sheet(src_sheet_name)
        dest_sheet = self._get_sheet(dest_sheet_name)

        for row in src_sheet.iter_rows():
            for cell in row:
                dest_sheet[cell.coordinate].value = cell.value

        self._log_audit("SUCCESS", f"Data copied from '{src_sheet_name}' to '{dest_sheet_name}'.", "copy_data")

    def _process_calculate_column_task(self, task):
        """Calculates a new column using an Excel formula ({row} = current row number)."""
        sheet_name = task.get("sheet")
        target_column_letter = task.get("target_column")
        start_row = task.get("start_row", 2)  # Assume header in row 1
        formula = task.get("formula")  # e.g., "=B{row}*C{row}"

        if not sheet_name or not target_column_letter or not formula:
            self._log_audit("ERROR", "Missing required fields for calculate_column task.", "calculate_column", task)
            return

        sheet = self._get_sheet(sheet_name)

        header_name = task.get("header", "Calculated Column")
        sheet[f"{target_column_letter}1"].value = header_name
        sheet[f"{target_column_letter}1"].font = Font(bold=True)

        for row_num in range(start_row, sheet.max_row + 1):
            cell_coord = f"{target_column_letter}{row_num}"
            try:
                sheet[cell_coord] = formula.replace("{row}", str(row_num))
                if task.get("number_format"):
                    sheet[cell_coord].number_format = task["number_format"]
            except Exception as e:
                self._log_audit("ERROR", f"Failed to apply formula to {cell_coord}: {e}", "calculate_column", task)

        self._log_audit("SUCCESS", f"Column '{target_column_letter}' calculated in '{sheet_name}'.", "calculate_column")

    def _process_conditional_format_task(self, task):
        """Applies conditional formatting to a range."""
        sheet_name = task.get("sheet")
        data_range = task.get("range")
        rule_type = task.get("rule_type")

        if not sheet_name or not data_range or not rule_type:
            self._log_audit("ERROR", "Missing required fields for conditional_format task.", "conditional_format", task)
            return

        sheet = self._get_sheet(sheet_name)

        try:
            if rule_type == "color_scale":
                # 3-color scale: ต่ำสุด -> กลาง -> สูงสุด
                rule = ColorScaleRule(start_type='min', start_color=Color(task.get("min_color", "FF0000")),
                                      mid_type='percentile', mid_value=50, mid_color=Color(task.get("mid_color", "FFFF00")),
                                      end_type='max', end_color=Color(task.get("max_color", "00FF00")))
            elif rule_type == "data_bar":
                rule = DataBarRule(start_type='min', end_type='max', color=Color(task.get("color", "0070C0")))
            elif rule_type == "icon_set":
                rule = IconSetRule(task.get("icon_set", '3TrafficLights1'), task.get("types", 'percent'),
                                   task.get("values", [0, 33, 67]))
            elif rule_type == "expression":
                # ไฮไลต์เซลล์เมื่อสูตรเป็นจริง
                fill_color = task.get("fill_color", "FFC7CE")
                dxf = DifferentialStyle(
                    font=Font(color=task.get("font_color", "9C0006")),
                    fill=PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid"))
                rule = Rule(type='expression', formula=[task.get("formula").lstrip("=")], dxf=dxf)
            else:
                self._log_audit("ERROR", f"Unsupported conditional format rule_type: {rule_type}", "conditional_format", task)
                return

            sheet.conditional_formatting.add(data_range, rule)
            self._log_audit("SUCCESS", f"Conditional formatting '{rule_type}' applied to '{data_range}' in '{sheet_name}'.",
                            "conditional_format")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to apply conditional formatting to '{data_range}': {e}", "conditional_format", task)

    def _process_chart_task(self, task):
        """Creates a chart based on specified data."""
        sheet_name = task.get("sheet")
        chart_type = task.get("chart_type")
        data_range = task.get("data_range")  # e.g., "D1:D5" (รวมหัวคอลัมน์)
        category_range = task.get("category_range")  # e.g., "A2:A5" for labels
        title = task.get("title", "Chart Title")
        top_left_cell = task.get("top_left_cell", "G2")  # Where to place the chart

        if not sheet_name or not chart_type or not data_range:
            self._log_audit("ERROR", "Missing required fields for chart task.", "chart", task)
            return

        sheet = self._get_sheet(sheet_name)

        try:
            chart_classes = {"bar": BarChart, "line": LineChart, "pie": PieChart}
            if chart_type not in chart_classes:
                self._log_audit("ERROR", f"Unsupported chart type: {chart_type}", "chart", task)
                return
            chart = chart_classes[chart_type]()

            min_col, min_row, max_col, max_row = range_boundaries(data_range)
            data = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.series.append(Series(data, title_from_data=True))  # แถวแรกของ data_range เป็นชื่อ series

            if category_range:
                min_col, min_row, max_col, max_row = range_boundaries(category_range)
                # ข้อมูลเริ่มแถวถัดจากหัวคอลัมน์ ดังนั้นป้ายกำกับก็ต้องเริ่มแถวเดียวกัน
                if min_row == 1 and range_boundaries(data_range)[1] == 1:
                    min_row = 2
                cats = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
                chart.set_categories(cats)

            chart.title = title
            if chart_type != "pie":  # pie chart ไม่มีแกน
                chart.x_axis.title = task.get("x_axis_title", "")
                chart.y_axis.title = task.get("y_axis_title", "")

            sheet.add_chart(chart, top_left_cell)
            self._log_audit("SUCCESS", f"Chart '{title}' of type '{chart_type}' added to '{sheet_name}'.", "chart")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to create chart '{title}': {e}", "chart", task)

    def _add_audit_log_sheet(self):
        """Adds a sheet to the workbook containing the audit log."""
        if not self.audit_log:
            return

        audit_sheet = self.workbook.create_sheet("Audit Log", 0)  # Create as first sheet
        audit_sheet.append(["Timestamp", "Status", "Task Type", "Message", "Details"])

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")  # Dark Gray
        for col_num in range(1, 6):
            cell = audit_sheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill

        for entry in self.audit_log:
            audit_sheet.append([
                entry["timestamp"],
                entry["status"],
                entry["task_type"],
                entry["message"],
                json.dumps(entry["details"], ensure_ascii=False)  # Store details as JSON string
            ])

        # Auto-adjust column widths
        for col_num in range(1, 6):
            column_letter = get_column_letter(col_num)
            max_length = max((len(str(cell.value)) for cell in audit_sheet[column_letter] if cell.value), default=0)
            if max_length > 0:
                audit_sheet.column_dimensions[column_letter].width = max_length + 2

        logger.info("Audit log sheet created.")

    def run(self):
        """Executes the defined tasks from the configuration."""
        if not self._load_workbook():
            self._log_audit("CRITICAL", "Agent could not start due to workbook loading failure.", "run")
            return

        handlers = {
            "copy_data": self._process_copy_data_task,
            "calculate_column": self._process_calculate_column_task,
            "conditional_format": self._process_conditional_format_task,
            "chart": self._process_chart_task,
        }

        try:
            for i, task in enumerate(self.config["tasks"]):
                task_type = task.get("type")
                logger.info(f"Executing Task {i + 1}: Type='{task_type}' Sheet='{task.get('sheet', task.get('dest_sheet'))}'")
                handler = handlers.get(task_type)
                if handler:
                    handler(task)
                else:
                    self._log_audit("WARNING", f"Unsupported task type encountered: {task_type}", "task_execution", task)

            duration = (datetime.datetime.now() - self.start_time).total_seconds()
            self._log_audit("INFO", f"All tasks executed in {duration:.2f} seconds.", "run_summary")

            # Add final audit log sheet, then save
            self._add_audit_log_sheet()

            output_file = self.config["output_file"]
            try:
                self.workbook.save(output_file)
                logger.info(f"Final workbook saved to: {output_file}")
            except Exception as e:
                logger.critical(f"Error saving output workbook '{output_file}': {e}")
        except Exception as e:
            self._log_audit("CRITICAL", f"An unhandled error occurred during agent execution: {e}", "run")
            logger.critical(f"An unhandled error occurred: {e}", exc_info=True)
        finally:
            duration = (datetime.datetime.now() - self.start_time).total_seconds()
            logger.info(f"Agent finished. Total time: {duration:.2f} seconds.")
