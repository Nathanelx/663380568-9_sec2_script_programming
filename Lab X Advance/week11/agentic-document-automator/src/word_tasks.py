# agentic-document-automator/src/word_tasks.py
import os

import docx

from utils import setup_logging

logger = setup_logging(__name__)

TABLE_PLACEHOLDER = "{{TABLE_PLACEHOLDER}}"


class WordTasks:
    """Handles advanced Word document operations for the Document Agent."""

    def __init__(self):
        logger.info("WordTasks initialized.")

    @staticmethod
    def _as_placeholder(key):
        """'REPORT_TITLE' -> '{{REPORT_TITLE}}' (ถ้าใส่ {{ }} มาแล้วก็ใช้ตามนั้น)"""
        return key if key.startswith("{{") else f"{{{{{key}}}}}"

    def _replace_placeholder_in_paragraph(self, paragraph, placeholder, value):
        """Replaces a placeholder in a paragraph's runs. Returns True if found."""
        if placeholder not in paragraph.text:
            return False
        for run in paragraph.runs:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, str(value))
        return True

    def _build_table(self, document, table_data):
        """สร้างตารางจาก list of lists (แถวแรกเป็นหัวตาราง ตัวหนา)"""
        cols = len(table_data[0])
        table = document.add_table(rows=0, cols=cols)
        try:
            table.style = 'Table Grid'
        except Exception:
            logger.warning("Style 'Table Grid' not available, using default table style.")
        for r_idx, row_data in enumerate(table_data):
            row = table.add_row()
            for c_idx, cell_data in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.text = str(cell_data)
                if r_idx == 0 and cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].bold = True
        return table

    def generate_report_from_template(self, template_path, output_path, data_placeholders, table_data=None):
        """
        Generates a Word document from a template.
        `data_placeholders` maps names (e.g. 'REPORT_TITLE') to values; the template uses '{{REPORT_TITLE}}'.
        `table_data` is a list of lists (first row = header), inserted at '{{TABLE_PLACEHOLDER}}'.
        """
        if not os.path.exists(template_path):
            logger.error(f"Word template not found: {template_path}")
            return False

        try:
            document = docx.Document(template_path)
            placeholders = {self._as_placeholder(k): v for k, v in data_placeholders.items()}

            # Replace placeholders in paragraphs and table cells
            paragraphs = list(document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            for paragraph in paragraphs:
                for placeholder, value in placeholders.items():
                    if self._replace_placeholder_in_paragraph(paragraph, placeholder, value):
                        logger.info(f"Replaced '{placeholder}' with '{value}'.")

            # Add dynamic table if provided
            if table_data:
                table = self._build_table(document, table_data)  # add_table ใส่ไว้ท้ายเอกสารก่อน
                anchor = next((p for p in document.paragraphs if TABLE_PLACEHOLDER in p.text), None)
                if anchor is not None:
                    # ย้ายตารางไปไว้ตำแหน่ง placeholder แล้วลบย่อหน้า placeholder ทิ้ง
                    anchor._element.addnext(table._tbl)
                    anchor._element.getparent().remove(anchor._element)
                    logger.info("Dynamic table inserted at '{{TABLE_PLACEHOLDER}}'.")
                else:
                    logger.warning("Table placeholder not found in template. Table appended to end.")

            document.save(output_path)
            logger.info(f"Report generated from template and saved to '{output_path}'.")
            return True

        except Exception as e:
            logger.error(f"Error generating report from template: {e}")
            return False

    def extract_text_from_word(self, docx_path):
        """Extracts all text (paragraphs + tables) from a Word document."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None

        try:
            document = docx.Document(docx_path)
            full_text = [para.text for para in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    full_text.append("\t".join(cell.text for cell in row.cells))  # Tab-separated cells
            logger.info(f"Text extracted from Word document: {docx_path}")
            return "\n".join(full_text)
        except Exception as e:
            logger.error(f"Error extracting text from Word document '{docx_path}': {e}")
            return None
